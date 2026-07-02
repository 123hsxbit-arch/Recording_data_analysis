from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def create_image_report(image_paths: list[Path], ppt_path: Path, title: str) -> Path:
    """Create a PPT report with one image per slide."""
    prs = Presentation()

    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"Total charts: {len(image_paths)}"

    blank_layout = prs.slide_layouts[6]
    for image_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image_path),
            Inches(0.5),
            Inches(0.5),
            width=Inches(9.0),
        )

    ppt_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(ppt_path)
    return ppt_path

